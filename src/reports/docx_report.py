"""
Word (.docx) equity research report generator.

Mirrors the PDF report's content so a user can take either format away, but as an
editable document — the point of offering Word is that a reader can annotate the
thesis, drop sections, and paste tables into their own notes.

Requires python-docx: pip install python-docx
"""
from __future__ import annotations

import os
from datetime import datetime
from typing import Any, Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.reports.disclosure import (
    AI_CONTENT_NOTICE,
    SEBI_DISCLAIMER,
    SEBI_DISCLAIMER_SHORT,
    provenance_from_result,
)
from src.reports.report_data import fmt_num, fmt_pct, load_report_financials

# Palette kept in step with the PDF report so the two formats look related.
_NAVY  = RGBColor(0x0F, 0x28, 0x50)
_GOLD  = RGBColor(0xB4, 0x8C, 0x28)
_GREY  = RGBColor(0x60, 0x60, 0x60)
_RED   = RGBColor(0xB0, 0x2A, 0x2A)
_GREEN = RGBColor(0x1A, 0x7F, 0x37)

_VERDICT_COLOURS = {
    "LIKELY UNDERVALUED": _GREEN,
    "LEANS UNDERVALUED":  _GREEN,
    "FAIRLY VALUED":      _GOLD,
    "LEANS OVERVALUED":   _RED,
    "LIKELY OVERVALUED":  _RED,
    # Composite-score verdicts
    "TOP PICK": _GREEN, "BUY": _GREEN, "HOLD": _GOLD, "AVOID": _RED,
}


def _shade(cell, hex_colour: str) -> None:
    """Apply a solid background fill to a table cell (python-docx has no API for this)."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shading)


def _run(paragraph, text: str, size: int = 10, bold: bool = False,
         italic: bool = False, color: Optional[RGBColor] = None):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _para(doc, text: str = "", size: int = 10, bold: bool = False,
          italic: bool = False, color: Optional[RGBColor] = None,
          align=None, space_after: int = 4):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    if text:
        _run(paragraph, text, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def _heading(doc, text: str, size: int = 13) -> None:
    paragraph = _para(doc, text, size=size, bold=True, color=_NAVY, space_after=2)
    # A bottom border under the heading, standing in for the PDF's rule line.
    borders = OxmlElement("w:pBdr")
    bottom  = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "0F2850")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _kv_table(doc, rows: list[tuple[str, str]], widths=(2.6, 3.4)) -> None:
    """Two-column label/value table used for the key-stats and valuation blocks."""
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        _shade(cells[0], "F2F4F8")
        p0 = cells[0].paragraphs[0]
        _run(p0, label, size=9, bold=True)
        p1 = cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(p1, value, size=9)


def _add_footer(section, text: str) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.text = ""
    _run(paragraph, text, size=7, italic=True, color=_GREY)


def generate_docx(
    symbol:       str,
    company_name: str,
    dcf_result:   Any,
    score_result: Any,
    ai_thesis:    str = "",
    output_dir:   str = "reports",
    cache_dir:    str = "data/raw/screener",
    freshness:    Optional[dict] = None,
    filename:     Optional[str] = None,
) -> str:
    """Generate the Word report and return the output file path."""
    os.makedirs(output_dir, exist_ok=True)
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    name     = filename or f"{symbol}_EquityResearch_{date_str}.docx"
    out_path = os.path.join(output_dir, name)

    fin = load_report_financials(symbol, cache_dir=cache_dir)
    doc = Document()

    # Page setup + a disclosure footer on every page, so a printed or forwarded
    # extract still carries it.
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.7)
    _add_footer(section, SEBI_DISCLAIMER_SHORT)

    # ── Cover ─────────────────────────────────────────────────────────────────
    _para(doc, "EQUITYLENS EQUITY RESEARCH", size=9, bold=True, color=_GOLD,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, company_name, size=22, bold=True, color=_NAVY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, f"NSE: {symbol}  |  {datetime.now().strftime('%d %B %Y')}",
          size=9, color=_GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    verdict = getattr(dcf_result, "verdict", "N/A")
    _para(doc, verdict, size=16, bold=True,
          color=_VERDICT_COLOURS.get(verdict, _NAVY),
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # ── Key stats ─────────────────────────────────────────────────────────────
    _heading(doc, "KEY STATISTICS")
    cmp_val = getattr(dcf_result, "cmp", None)
    iv      = getattr(dcf_result, "intrinsic_value", 0) or 0
    mos     = getattr(dcf_result, "margin_of_safety", None)

    stats: list[tuple[str, str]] = [
        ("Current market price", f"Rs {fmt_num(cmp_val, 2)}" if cmp_val else "—"),
        ("Intrinsic value (model)", f"Rs {fmt_num(iv, 2)}" if iv else "—"),
        ("Valuation range (P10–P90)",
         f"Rs {fmt_num(getattr(dcf_result, 'iv_low', 0), 0)} – "
         f"Rs {fmt_num(getattr(dcf_result, 'iv_high', 0), 0)}"),
        ("Margin of safety", f"{mos*100:+.1f}%" if mos is not None else "—"),
        ("Valuation model", getattr(dcf_result, "model", "—")),
        ("Discount rate (WACC / CoE)", fmt_pct(getattr(dcf_result, "wacc", 0) * 100, 2)),
    ]
    if score_result is not None:
        stats += [
            ("Composite quality score",
             f"{getattr(score_result, 'composite', 0):.1f} / 100 "
             f"(grade {getattr(score_result, 'grade', '—')})"),
            ("Piotroski F-Score", f"{getattr(score_result, 'piotroski', 0)} / 9"),
            ("Revenue CAGR (3Y)", fmt_pct(getattr(score_result, "revenue_cagr_3y", 0))),
            ("EPS CAGR (3Y)", fmt_pct(getattr(score_result, "eps_cagr_3y", 0))),
            ("Debt / Equity", f"{getattr(score_result, 'de_ratio', 0):.2f}x"),
        ]
    _kv_table(doc, stats)

    # ── Five-year financials ──────────────────────────────────────────────────
    if fin.has_data:
        doc.add_paragraph()
        _heading(doc, "FINANCIAL SUMMARY — 5-YEAR TREND")
        _para(doc, "Figures in Rs. Crore unless marked otherwise.",
              size=8, italic=True, color=_GREY, space_after=4)

        metrics = [
            ("Revenue" if fin.is_bank else "Sales", fin.sales, 0, False),
            ("Net profit",                          fin.net_profit, 0, False),
            ("EPS (Rs)",                            fin.eps, 2, False),
            (fin.opm_label,                         fin.opm, 1, True),
            (fin.roce_label,                        fin.roce, 1, True),
            ("Cash from operations",                fin.cfo, 0, False),
            ("Free cash flow",                      fin.fcf, 0, False),
        ]

        table = doc.add_table(rows=1, cols=len(fin.years) + 1)
        table.style = "Table Grid"
        header = table.rows[0].cells
        _shade(header[0], "0F2850")
        _run(header[0].paragraphs[0], "Metric", size=8, bold=True,
             color=RGBColor(0xFF, 0xFF, 0xFF))
        for i, year in enumerate(fin.years, start=1):
            _shade(header[i], "0F2850")
            header[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(header[i].paragraphs[0], year, size=8, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF))

        for row_idx, (label, values, decimals, is_pct) in enumerate(metrics):
            cells = table.add_row().cells
            if row_idx % 2 == 0:
                for cell in cells:
                    _shade(cell, "F7F8FA")
            _run(cells[0].paragraphs[0], label, size=8, bold=True)
            for i, value in enumerate(values, start=1):
                if i >= len(cells):
                    break
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                text = fmt_pct(value, decimals) if is_pct else fmt_num(value, decimals)
                _run(cells[i].paragraphs[0], text, size=8)

    # ── Valuation detail ──────────────────────────────────────────────────────
    doc.add_paragraph()
    _heading(doc, "VALUATION")
    model = getattr(dcf_result, "model", "DCF")
    model_notes = {
        "DCF": "Free-cash-flow DCF: 10-year explicit projection plus a Gordon-growth "
               "terminal value, discounted at WACC. 10,000 Monte Carlo draws give the "
               "P10–P90 range.",
        "RIM": "Residual Income Model, used for banks and NBFCs because their cash "
               "flow statements are not comparable to an operating company's. "
               "IV = book value + present value of excess returns over cost of equity.",
        "PE":  "Relative P/E model, used for life insurers because reported operating "
               "cash flow includes policyholder premiums and makes a DCF meaningless.",
    }
    _para(doc, model_notes.get(model, ""), size=9, space_after=6)

    _kv_table(doc, [
        ("Model", model),
        ("PV of explicit forecast", f"Rs {fmt_num(getattr(dcf_result, 'pv_fcf', 0))} Cr"),
        ("PV of terminal value", f"Rs {fmt_num(getattr(dcf_result, 'terminal_value', 0))} Cr"),
        ("Shares outstanding", f"{fmt_num(getattr(dcf_result, 'shares_cr', 0), 2)} Cr"),
        ("Growth input",
         fmt_pct(getattr(dcf_result, "fcf_cagr", 0) * 100, 1)
         if model != "PE" else f"Rs {fmt_num(getattr(dcf_result, 'fcf_cagr', 0), 2)} avg EPS"),
    ])

    notes = getattr(dcf_result, "notes", None) or []
    if notes:
        doc.add_paragraph()
        _para(doc, "Model notes", size=10, bold=True, space_after=2)
        for note in notes:
            bullet = doc.add_paragraph(style="List Bullet")
            _run(bullet, str(note), size=9)

    # ── AI thesis ─────────────────────────────────────────────────────────────
    if ai_thesis:
        doc.add_paragraph()
        _heading(doc, "INVESTMENT THESIS")
        for block in [b.strip() for b in ai_thesis.split("\n\n") if b.strip()]:
            _para(doc, block, size=10, space_after=6)
        _para(doc, AI_CONTENT_NOTICE, size=8, italic=True, color=_GREY, space_after=4)

    # ── Provenance + disclosure ───────────────────────────────────────────────
    doc.add_paragraph()
    _heading(doc, "DATA PROVENANCE")
    for line in provenance_from_result(dcf_result, freshness):
        bullet = doc.add_paragraph(style="List Bullet")
        _run(bullet, line, size=8,
             color=_RED if ("STALE" in line or "NOT LIVE" in line) else _GREY)

    doc.add_paragraph()
    _heading(doc, "IMPORTANT DISCLOSURES")
    _para(doc, SEBI_DISCLAIMER, size=8, italic=True, color=_GREY, space_after=4)

    doc.save(out_path)
    return out_path
